from __future__ import annotations

import openpyxl
import pandas as pd

from . import io as rio
from .build_correspondences_abs import _canonical_label_lookup, _lexical_score
from .hierarchy import BRIDGE_COLUMNS, write_csv
from .paths import BRIDGES_DIR, RAW_DIR

# Already .xlsx (user-converted from the original .xls), so no LibreOffice conversion step
# is needed here anymore.
LEGACY_XLSX = RAW_DIR / "abs_for_seo" / "1297.0 correspondence tables.xlsx"


def parse_1998_to_2008(sheet: str) -> pd.DataFrame:
    """Table 1 (RFCD1998, called FOR1998 in this project's own naming -> FOR2008) /
    Table 3 (SEO1998 -> SEO2008): plain code/label columns, with an optional trailing 'p'
    glued onto the 2008 code for partial matches."""
    wb = openpyxl.load_workbook(LEGACY_XLSX, data_only=True)
    ws = wb[sheet]
    rows = []
    for row in ws.iter_rows(min_row=6, max_row=ws.max_row, max_col=4, values_only=True):
        code_1998, name_1998, code_2008_raw, name_2008 = row
        if code_1998 is None or code_2008_raw is None:
            continue
        parsed = rio.split_trailing_p_code(code_2008_raw)
        if parsed is None:
            continue
        code_2008, is_partial = parsed
        rows.append(
            {
                "source_code": str(int(code_1998)) if isinstance(code_1998, (int, float)) else str(code_1998).strip(),
                "source_label": str(name_1998).strip(),
                "code_2008": code_2008,
                "label_2008": str(name_2008).strip(),
                "is_partial": is_partial,
            }
        )
    return pd.DataFrame(rows)


def resolve_1998_to_2008_primary(raw: pd.DataFrame) -> pd.DataFrame:
    """Same single-vs-multi-candidate resolution rule as the 2008<->2020 correspondence:
    one candidate -> explicit_official (confidence 1.0, or the lexical score if p-flagged);
    multiple candidates -> lexical tiebreak among them, one primary."""
    out_rows = []
    for source_code, grp in raw.groupby("source_code"):
        source_label = grp["source_label"].iloc[0]
        if len(grp) == 1:
            r = grp.iloc[0]
            confidence = 1.0 if not r["is_partial"] else _lexical_score(source_label, r["label_2008"])
            out_rows.append(
                {
                    "source_code": source_code,
                    "source_label": source_label,
                    "code_2008": r["code_2008"],
                    "label_2008": r["label_2008"],
                    "is_primary": True,
                    "confidence": round(confidence, 3),
                }
            )
        else:
            scored = grp.copy()
            scored["score"] = scored["label_2008"].apply(lambda lbl: _lexical_score(source_label, lbl))
            scored = scored.sort_values(["score", "code_2008"], ascending=[False, True])
            for i, (_, r) in enumerate(scored.iterrows()):
                out_rows.append(
                    {
                        "source_code": source_code,
                        "source_label": source_label,
                        "code_2008": r["code_2008"],
                        "label_2008": r["label_2008"],
                        "is_primary": i == 0,
                        "confidence": round(r["score"], 3),
                    }
                )
    return pd.DataFrame(out_rows)


def compose_with_2008_2020(
    resolved_1998: pd.DataFrame, bridge_2008_2020: pd.DataFrame, system: str, source_system: str
) -> pd.DataFrame:
    """Second hop is deterministic pass-through: bridge_2008_2020 already has exactly one
    primary per 2008 code, so composing doesn't introduce extra branching -- each 1998->2008
    candidate (primary or not) becomes one 1998->2020 row, keeping its own is_primary flag."""
    lookup_2020 = bridge_2008_2020[bridge_2008_2020["is_primary"] == True].set_index("source_code")  # noqa: E712
    canonical_lookup = _canonical_label_lookup(system)

    rows = []
    unresolved = 0
    for _, r in resolved_1998.iterrows():
        if r["code_2008"] not in lookup_2020.index:
            unresolved += 1
            continue
        target = lookup_2020.loc[r["code_2008"]]
        rows.append(
            {
                "source_system": source_system,
                "source_code": r["source_code"],
                "source_label": r["source_label"],
                "system": system,
                "canonical_code": target["canonical_code"],
                "canonical_label": canonical_lookup.get(target["canonical_code"], target["canonical_label"]),
                "canonical_level": target["canonical_level"],
                "is_primary": r["is_primary"],
                "match_method": "explicit_official_transitive",
                "confidence": round(min(float(r["confidence"]), 0.9), 3),
                "notes": f"{source_system}->{r['code_2008']} ({'2008'})->{target['canonical_code']} (2020)",
            }
        )
    if unresolved:
        print(f"  [{source_system}] {unresolved} row(s) had no matching primary in the 2008->2020 bridge, dropped")
    return pd.DataFrame(rows, columns=BRIDGE_COLUMNS)


def run() -> dict[str, pd.DataFrame]:
    for1998_raw = parse_1998_to_2008("Table 1")
    seo1998_raw = parse_1998_to_2008("Table 3")

    for1998_resolved = resolve_1998_to_2008_primary(for1998_raw)
    seo1998_resolved = resolve_1998_to_2008_primary(seo1998_raw)

    for2008_2020 = pd.read_csv(BRIDGES_DIR / "bridge_for2008_for2020.csv", dtype=str, keep_default_na=False)
    for2008_2020["is_primary"] = for2008_2020["is_primary"].isin(["True", "true"])
    seo2008_2020 = pd.read_csv(BRIDGES_DIR / "bridge_seo2008_seo2020.csv", dtype=str, keep_default_na=False)
    seo2008_2020["is_primary"] = seo2008_2020["is_primary"].isin(["True", "true"])

    # "FOR1998" is this project's own naming for consistency with FOR2008/FOR2020; ABS's own
    # documents call this scheme RFCD1998 (Research Fields, Courses and Disciplines).
    for1998_bridge = compose_with_2008_2020(for1998_resolved, for2008_2020, "FOR", "FOR1998")
    seo1998_bridge = compose_with_2008_2020(seo1998_resolved, seo2008_2020, "SEO", "SEO1998")

    write_csv(for1998_bridge, BRIDGES_DIR / "bridge_for1998_for2020.csv", ["source_code"])
    write_csv(seo1998_bridge, BRIDGES_DIR / "bridge_seo1998_seo2020.csv", ["source_code"])

    return {
        "bridge_for1998_for2020": for1998_bridge,
        "bridge_seo1998_seo2020": seo1998_bridge,
    }


if __name__ == "__main__":
    tables = run()
    for name, df in tables.items():
        print(name, len(df), "primaries:", df["is_primary"].sum())
