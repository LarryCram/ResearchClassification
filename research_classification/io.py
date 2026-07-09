"""Shared parsing utilities: legacy .xls conversion and ANZSRC-style hierarchy readers."""

from __future__ import annotations

import subprocess
from pathlib import Path

import openpyxl
import pandas as pd

RAW_DIR = Path(__file__).resolve().parent.parent / "data_untracked"
CONVERTED_DIR = RAW_DIR / "_converted"


def ensure_converted(xls_path: Path) -> Path:
    """Convert a legacy .xls file to .xlsx via headless LibreOffice, caching the result.

    Re-converts only if the source is newer than any previously converted copy.
    """
    CONVERTED_DIR.mkdir(parents=True, exist_ok=True)
    target = CONVERTED_DIR / (xls_path.stem + ".xlsx")
    if target.exists() and target.stat().st_mtime >= xls_path.stat().st_mtime:
        return target
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(CONVERTED_DIR),
            str(xls_path),
        ],
        check=True,
        capture_output=True,
    )
    if not target.exists():
        raise RuntimeError(f"soffice conversion did not produce {target}")
    return target


def read_indented_hierarchy(
    xlsx_path: Path,
    sheet_name: str,
    n_levels: int,
    level_names: list[str],
    header_row: int,
) -> pd.DataFrame:
    """Parse ABS-style ANZSRC hierarchy sheets where each level's code lives in its own
    column (col 0 = level-1 code, col 1 = level-2 code, ...) with the label in the column
    immediately to its right, and the level is determined by which column holds an int code.

    Returns tidy long form: code, level, label, parent_code (all str; parent_code="" for roots).
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    rows: list[dict] = []
    current: dict[int, str] = {}
    for row in ws.iter_rows(
        min_row=header_row + 1, max_row=ws.max_row, max_col=n_levels + 1, values_only=True
    ):
        level_idx = None
        for i in range(n_levels):
            if isinstance(row[i], int):
                level_idx = i
                break
        if level_idx is None:
            continue
        code = str(row[level_idx])
        label = row[level_idx + 1]
        if label is None:
            continue
        label = str(label).strip()
        current[level_idx] = code
        for deeper in range(level_idx + 1, n_levels):
            current.pop(deeper, None)
        parent_code = current.get(level_idx - 1, "") if level_idx > 0 else ""
        rows.append(
            {
                "code": code,
                "level": level_names[level_idx],
                "label": label,
                "parent_code": parent_code,
            }
        )
    return pd.DataFrame(rows)


def read_definitions(xlsx_path: Path, sheet_name: str = "Table 4") -> pd.DataFrame:
    """Parse the ABS 'Table 4' definitions/exclusions sheet (division + group level only).

    Returns columns: code, level, label, definition, exclusions.
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet_name]
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=5, values_only=True):
        if isinstance(row[0], int):
            code, label, level = str(row[0]), row[1], "division"
            definition, exclusions = row[3], row[4]
        elif isinstance(row[1], int):
            code, label, level = str(row[1]), row[2], "group"
            definition, exclusions = row[3], row[4]
        else:
            continue
        if label is None:
            continue
        rows.append(
            {
                "code": code,
                "level": level,
                "label": str(label).strip(),
                "definition": (definition or "").strip() if definition else "",
                "exclusions": (exclusions or "").strip() if exclusions else "",
            }
        )
    return pd.DataFrame(rows)


def split_trailing_p_code(cell: object) -> tuple[str, bool] | None:
    """Split a code with an optional trailing 'p' glued directly onto it (used by the
    legacy '1297.0 correspondence tables.xls' RFCD1998<->FOR2008 / SEO1998<->SEO2008
    tables, e.g. '010104p') into (code, is_partial).
    """
    if cell is None:
        return None
    text = str(cell).strip()
    if not text:
        return None
    is_partial = text.endswith(("p", "P"))
    code = text[:-1] if is_partial else text
    if not code.isdigit():
        return None
    return code, is_partial
