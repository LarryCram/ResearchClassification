from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from .hierarchy import write_csv

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "research_classification" / "data"
SRC = DATA_DIR / "OpenAlex_topic_mapping_table.xlsx"  # tracked in git, not data_untracked --
# this is the pipeline's actual raw source, checked in directly so the build doesn't depend
# on a gitignored local copy for this one file.


def _to_int_str(v: object) -> str:
    return str(int(float(v)))


def _fix_mojibake(text: str) -> str:
    """A minority of rows in the source xlsx (~1-2%) were corrupted upstream: correct UTF-8
    text got decoded as Mac OS Roman at some point before being re-saved as UTF-8, e.g.
    'Catalytic C‚ÄìH...' should read 'Catalytic C–H...'. Reversing that (encode as the
    encoding it was wrongly read as, decode as the encoding it should have been read as)
    recovers the original text. Genuinely correct UTF-8 elsewhere (e.g. Māori macrons
    elsewhere in the pipeline) isn't representable in Mac OS Roman, so this safely no-ops
    (UnicodeEncodeError) rather than risking corrupting text that was never broken."""
    try:
        fixed = text.encode("mac_roman").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return text
    return fixed


def load_raw() -> pd.DataFrame:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {name: i for i, name in enumerate(header) if name}
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[col_idx["topic_id"]] is None:
            continue
        rows.append(
            {
                "topic_id": _to_int_str(row[col_idx["topic_id"]]),
                "topic_name": _fix_mojibake(str(row[col_idx["topic_name"]]).strip()),
                "subfield_id": _to_int_str(row[col_idx["subfield_id"]]),
                "subfield_name": _fix_mojibake(str(row[col_idx["subfield_name"]]).strip()),
                "field_id": _to_int_str(row[col_idx["field_id"]]),
                "field_name": _fix_mojibake(str(row[col_idx["field_name"]]).strip()),
                "domain_id": _to_int_str(row[col_idx["domain_id"]]),
                "domain_name": _fix_mojibake(str(row[col_idx["domain_name"]]).strip()),
                "keywords": _fix_mojibake(row[col_idx["keywords"]] or ""),
                "summary": _fix_mojibake(row[col_idx["summary"]] or ""),
                "wikipedia_url": (row[col_idx["wikipedia_url"]] or "").strip().rstrip("/"),
            }
        )
    return pd.DataFrame(rows)


def run() -> dict[str, pd.DataFrame]:
    raw = load_raw()

    domains = raw[["domain_id", "domain_name"]].drop_duplicates().rename(
        columns={"domain_id": "code", "domain_name": "label"}
    )
    domains["level"] = "domain"
    domains["parent_code"] = ""

    fields = raw[["field_id", "field_name", "domain_id"]].drop_duplicates().rename(
        columns={"field_id": "code", "field_name": "label", "domain_id": "parent_code"}
    )
    fields["level"] = "field"

    subfields = raw[["subfield_id", "subfield_name", "field_id"]].drop_duplicates().rename(
        columns={"subfield_id": "code", "subfield_name": "label", "field_id": "parent_code"}
    )
    subfields["level"] = "subfield"

    topics = raw[["topic_id", "topic_name", "subfield_id"]].drop_duplicates().rename(
        columns={"topic_id": "code", "topic_name": "label", "subfield_id": "parent_code"}
    )
    topics["level"] = "topic"

    cols = ["code", "level", "label", "parent_code"]
    tables = {
        "openalex_domains": domains[cols],
        "openalex_fields": fields[cols],
        "openalex_subfields": subfields[cols],
        "openalex_topics": topics[cols],
    }
    for name, df in tables.items():
        write_csv(df, DATA_DIR / f"{name}.csv", ["code"])

    # keep the enriched topic-level table (with keywords/summary/wikipedia_url) available
    # in-memory for the Leiden join and the curation self-check; not written as a separate
    # canonical file since 'level' semantics already live in openalex_topics.csv.
    return {**tables, "_raw": raw}


if __name__ == "__main__":
    tables = run()
    for name, df in tables.items():
        print(name, len(df))
