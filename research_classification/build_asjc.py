from __future__ import annotations

from pathlib import Path

import openpyxl
import pandas as pd

from .hierarchy import BRIDGE_COLUMNS, write_csv

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data_untracked" / "ASJC1.xlsx"
DATA_DIR = ROOT / "research_classification" / "data"


def load_asjc() -> pd.DataFrame:
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["ASJC"]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=3, values_only=True):
        code_system, code, label = row
        if code is None or label is None:
            continue
        code = str(int(code))
        level = "field" if len(code) == 2 else "subfield"
        rows.append({"code": code, "level": level, "label": str(label).strip()})
    return pd.DataFrame(rows).drop_duplicates(subset=["code"])


def build_bridge(asjc: pd.DataFrame, openalex_fields: pd.DataFrame, openalex_subfields: pd.DataFrame) -> pd.DataFrame:
    oax_codes = {
        **{c: (l, "field") for c, l in zip(openalex_fields["code"], openalex_fields["label"])},
        **{c: (l, "subfield") for c, l in zip(openalex_subfields["code"], openalex_subfields["label"])},
    }
    rows = []
    for _, r in asjc.iterrows():
        match = oax_codes.get(r["code"])
        if match is None:
            continue
        oax_label, oax_level = match
        rows.append(
            {
                "source_system": "ASJC",
                "source_code": r["code"],
                "source_label": r["label"],
                "system": "OAX",
                "canonical_code": r["code"],
                "canonical_label": oax_label,
                "canonical_level": oax_level,
                "is_primary": True,
                "match_method": "exact_key_join",
                "confidence": 1.0,
                "notes": "ASJC code == OpenAlex field_id/subfield_id (verified 100% coverage)",
            }
        )
    return pd.DataFrame(rows, columns=BRIDGE_COLUMNS)


def run() -> tuple[pd.DataFrame, pd.DataFrame]:
    asjc = load_asjc()
    openalex_fields = pd.read_csv(DATA_DIR / "openalex_fields.csv", dtype=str, keep_default_na=False)
    openalex_subfields = pd.read_csv(DATA_DIR / "openalex_subfields.csv", dtype=str, keep_default_na=False)
    bridge = build_bridge(asjc, openalex_fields, openalex_subfields)

    write_csv(asjc, DATA_DIR / "asjc.csv", ["code"])
    write_csv(bridge, DATA_DIR / "bridge_asjc_openalex.csv", ["source_code"])
    return asjc, bridge


if __name__ == "__main__":
    asjc, bridge = run()
    print("asjc rows:", len(asjc))
    print("bridge rows:", len(bridge), "/ unmatched:", len(asjc) - len(bridge))
